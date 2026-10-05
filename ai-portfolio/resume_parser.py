import io
from pypdf import PdfReader
from pptx import Presentation


def extract_text(filename: str, data: bytes) -> str:
    """Extract readable text from uploaded PDF, PPTX, or TXT resume files."""
    name = filename.lower()

    try:
        if name.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("Password-protected PDFs are not supported.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            return text.strip()

        if name.endswith(".pptx"):
            prs = Presentation(io.BytesIO(data))
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        texts.append(shape.text_frame.text)
            return "\n".join(texts).strip()

        if name.endswith(".txt"):
            return data.decode("utf-8", errors="replace").strip()

    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Failed to read file '{filename}': file may be corrupted or malformed ({e})") from e

    raise ValueError("Unsupported file format. Please upload a PDF, PPTX, or TXT file.")